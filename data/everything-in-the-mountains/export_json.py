"""
Export Everything in the Mountains Excel workbook to a clean, structured JSON file.
Reads Alpine_Wave_Synthetic_Mountain_Data.xlsx (or everything-in-the-mountains.xlsx)
and produces data/everything-in-the-mountains/everything-in-the-mountains.json.
"""

import os
import sys
import json
import datetime
from openpyxl import load_workbook

def serialize_val(val):
    if val is None:
        return None
    if isinstance(val, (datetime.date, datetime.datetime)):
        return val.strftime("%Y-%m-%d")
    return val

def export_mountain_data_to_json(xlsx_path=None, json_path=None):
    base_dir = os.path.dirname(os.path.abspath(__file__))

    if xlsx_path is None:
        primary = os.path.join(base_dir, "Alpine_Wave_Synthetic_Mountain_Data.xlsx")
        fallback = os.path.join(base_dir, "everything-in-the-mountains.xlsx")
        xlsx_path = primary if os.path.exists(primary) else fallback

    if json_path is None:
        json_path = os.path.join(base_dir, "everything-in-the-mountains.json")

    if not os.path.exists(xlsx_path):
        raise FileNotFoundError(f"Workbook not found at {xlsx_path}")

    wb = load_workbook(filename=xlsx_path, data_only=True)

    result = {
        "metadata": {
            "shop_name": "Everything in the Mountains",
            "data_classification": "SYNTHETIC DATA: made-up showcase shop, not a real business",
            "phone": "SECOND NUMBER (not yet set up)",
            "hours": "8:00 AM – 6:00 PM daily in season, stat holidays included",
            "location": "Village Slopeside & West Alley, Okanagan, British Columbia",
            "seasons": ["Winter", "Summer", "Shoulder / All Year"],
            "source_workbook": os.path.basename(xlsx_path),
            "last_updated": datetime.date.today().strftime("%Y-%m-%d")
        },
        "sheets": {}
    }

    print(f"Loading {xlsx_path}...")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            result["sheets"][sheet_name] = []
            continue

        # Check for title/subtitle and header row
        # In our format: row 1 = title, row 2 = subtitle, row 3 = blank, row 4 = headers, row 5+ = data
        if len(rows) >= 4 and rows[3] and any(rows[3]):
            header_row_idx = 3
        else:
            # Fallback if headers start on row 1
            header_row_idx = 0

        raw_headers = rows[header_row_idx]
        headers = []
        for i, h in enumerate(raw_headers):
            if h is not None and str(h).strip():
                headers.append(str(h).strip())
            else:
                headers.append(f"Column_{i+1}")

        sheet_records = []
        for row in rows[header_row_idx + 1:]:
            # Skip empty rows
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue

            record = {}
            for col_idx, key in enumerate(headers):
                val = row[col_idx] if col_idx < len(row) else None
                record[key] = serialize_val(val)
            sheet_records.append(record)

        result["sheets"][sheet_name] = sheet_records
        print(f"  - Sheet '{sheet_name}': {len(sheet_records)} rows exported")

    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(result, f, indent=2, ensure_ascii=False)

    print(f"\nSuccessfully wrote JSON to {json_path}")
    print(f"Total sheets: {len(result['sheets'])}")
    return json_path

if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else None
    dest = sys.argv[2] if len(sys.argv) > 2 else None
    export_mountain_data_to_json(src, dest)
